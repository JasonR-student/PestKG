"""Build an EDITABLE PPTX for figure_02 from the real released data.

Every node, edge, bar, dot and label is a NATIVE PowerPoint shape (rounded
rectangle / connector / oval / text box), so each can be selected, moved,
coloured and re-typeset in PowerPoint -- no flattened image.

Layout coordinates mirror rebuild_figure02_realdata.py (x:0..13, y:0..7).
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SAMPLE = PROJECT_ROOT / "data" / "releases" / "2026.08.3_federated" / "sample"
OUTPUT = PROJECT_ROOT / "artifacts" / "research" / "figures" / "output"
FORTIN_ID = "AU:PRODUCT:7feed95ebf8a4d6f19180bcc"

SLIDE_W, SLIDE_H = 13.333, 7.5
X0, Y0 = 13.0, 7.0  # data-coordinate span

INK = RGBColor(0x17, 0x21, 0x2B)
MUTED = RGBColor(0x62, 0x70, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF4, 0xF7, 0xF8)
EDGE_GRAY = RGBColor(0x9A, 0xA7, 0xB2)
EDGE_LIGHT = RGBColor(0xB9, 0xC2, 0xCB)

C_SOURCE = RGBColor(0x6F, 0x81, 0x92)
C_COUNTRY = RGBColor(0x4E, 0x8D, 0x7C)
C_REG = RGBColor(0x78, 0x67, 0xA8)
C_PRODUCT = RGBColor(0xD8, 0x89, 0x45)
C_USE = RGBColor(0xC8, 0x5C, 0x5C)
C_AI = RGBColor(0x2C, 0x8E, 0x9E)

CAT_COLORS = [RGBColor(0x3E, 0x7C, 0x4F), RGBColor(0x86, 0xA8, 0x4E),
              RGBColor(0xC2, 0xA8, 0x5E), RGBColor(0x5F, 0x8F, 0xA3)]
CAT_NAMES = ["Perennial fruit & nut crops", "Field & row crops",
             "Non-agricultural & land areas", "Application methods"]


def ix(x: float) -> Inches:
    return Inches(x / X0 * SLIDE_W)


def iy(y: float) -> Inches:
    return Inches((Y0 - y) / Y0 * SLIDE_H)


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def neighborhood(product_id):
    nodes = read_csv(SAMPLE / "nodes.csv")
    edges = read_csv(SAMPLE / "edges.csv")
    by_id = {n["id"]: n for n in nodes}
    seen, frontier = {product_id}, {product_id}
    kept = []
    for _ in range(3):
        nxt = set()
        for e in edges:
            if e["start_id"] in frontier or e["end_id"] in frontier:
                kept.append(e)
                for nid in (e["start_id"], e["end_id"]):
                    if nid not in seen:
                        seen.add(nid)
                        nxt.add(nid)
        frontier = nxt
    return by_id, kept


def classify(label: str) -> int:
    u = label.upper()
    if "OVER 3 YEARS OLD" in u or "VINEYARD" in u:
        return 0
    if any(k in u for k in ("COTTON", "CHICKPEA", "BEAN", "PEANUT", "SOYBEAN",
                            "ONION", "SUGAR CANE", "PASTURE")):
        return 1
    if any(k in u for k in ("FOLIAR", "STEM INJECTION", "CUT STUMP")):
        return 3
    return 2


# ---------- shape helpers ----------
def set_arrowhead(connector):
    ln = connector.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)


def set_dashed(connector):
    ln = connector.line._get_or_add_ln()
    d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
    ln.append(d)


def add_text(slide, x, y, w, h, runs, *, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             fill=None):
    """runs: list of paragraphs; each paragraph = list of (text,size,bold,color)."""
    box = slide.shapes.add_textbox(ix(x), iy(y + h), Inches(w / X0 * SLIDE_W), Inches(h / Y0 * SLIDE_H))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(1)
    tf.margin_top = tf.margin_bottom = Pt(1)
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        box.line.fill.background()
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for (text, size, bold, color) in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "Arial"
    return box


def add_node_box(slide, cx, cy, w, h, title, subs, color):
    left, top = cx - w / 2, cy + h / 2
    # rounded box
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 ix(left), iy(top), Inches(w / X0 * SLIDE_W), Inches(h / Y0 * SLIDE_H))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = color
    box.line.width = Pt(1.4)
    box.shadow.inherit = False
    # top colour band
    band_h = 0.16
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  ix(left), iy(top), Inches(w / X0 * SLIDE_W), Inches(band_h / Y0 * SLIDE_H))
    band.fill.solid()
    band.fill.fore_color.rgb = color
    band.line.fill.background()
    band.shadow.inherit = False
    # text
    runs = [[(title, 12.5, True, INK)]]
    for s in subs:
        runs.append([(s, 9.5, False, MUTED)])
    add_text(slide, left, cy - h / 2, w, h, runs)


def border_pt(cx, cy, hw, hh, tx, ty):
    dx, dy = tx - cx, ty - cy
    if dx == 0 and dy == 0:
        return cx, cy
    sx = hw / abs(dx) if dx else float("inf")
    sy = hh / abs(dy) if dy else float("inf")
    s = min(sx, sy)
    return cx + dx * s, cy + dy * s


def add_edge(slide, a, b, ha, hb, label=None, label_xy=None, dashed=False, light=False):
    """a,b = (cx,cy); ha,hb = (half_w,half_h) of the two boxes."""
    p1 = border_pt(*a, *ha, *b)
    p2 = border_pt(*b, *hb, *a)
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      ix(p1[0]), iy(p1[1]), ix(p2[0]), iy(p2[1]))
    conn.line.color.rgb = EDGE_LIGHT if light else EDGE_GRAY
    conn.line.width = Pt(0.9 if light else 1.2)
    if dashed:
        set_dashed(conn)
    set_arrowhead(conn)
    if label:
        lx, ly = label_xy if label_xy else ((a[0] + b[0]) / 2, (a[1] + b[1]) / 2)
        tw, th = 1.7, 0.34
        add_text(slide, lx - tw / 2, ly - th / 2, tw, th,
                 [[(label, 9, False, MUTED)]], fill=WHITE)


def main():
    by_id, edges = neighborhood(FORTIN_ID)
    reg = next(n for n in by_id.values() if n["type"] == "Registration" and n["label_en"] == "84962")
    country = next(n for n in by_id.values() if n["type"] == "CountryJurisdiction")
    product = by_id[FORTIN_ID]
    use = next(n for n in by_id.values() if n["type"] == "RegistrationUse")
    ai = next(by_id[e["end_id"]] for e in edges
              if e["predicate"] == "hasActiveIngredient" and e["start_id"] == use["id"])
    crops = sorted({by_id[e["end_id"]]["label_en"] for e in edges
                    if e["predicate"] == "registeredForCrop" and e["start_id"] == use["id"]})
    props = json.loads(reg.get("properties_json") or "{}")
    reg56 = next(n for n in by_id.values() if n["type"] == "Registration" and n["label_en"] == "56708")
    props56 = json.loads(reg56.get("properties_json") or "{}")

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    # title
    add_text(slide, 0.25, 6.55, 11.5, 0.35,
             [[("One official website to a jurisdictional knowledge graph", 16, True, INK)]],
             align=PP_ALIGN.LEFT)
    add_text(slide, 0.02, 6.6, 0.4, 0.35, [[("a", 16, True, INK)]], align=PP_ALIGN.CENTER)

    # ---------- edges first (under boxes) ----------
    apvma, aus = (1.75, 6.0), (4.55, 6.0)
    regc, reg56c = (4.55, 4.35), (2.15, 4.35)
    fortin, use_c, ai_c = (7.35, 5.25), (7.35, 3.85), (7.35, 2.25)
    hw = dict(apvma=(1.075, 0.39), aus=(0.9, 0.39), regc=(1.175, 0.5), reg56c=(1.0, 0.475),
              fortin=(1.0, 0.39), use_c=(1.0, 0.475), ai_c=(1.0, 0.39))

    add_edge(slide, apvma, aus, hw["apvma"], hw["aus"], "source provenance", (3.15, 6.0), dashed=True)
    add_edge(slide, aus, regc, hw["aus"], hw["regc"], "hasRegistration", (4.55, 5.18))
    add_edge(slide, aus, reg56c, hw["aus"], hw["reg56c"], dashed=True, light=True)
    add_edge(slide, regc, fortin, hw["regc"], hw["fortin"], "hasProduct", (5.95, 4.95))
    add_edge(slide, regc, use_c, hw["regc"], hw["use_c"], "hasRegistrationUse", (5.95, 3.95))
    add_edge(slide, use_c, fortin, hw["use_c"], hw["fortin"], "usesProduct", (7.35, 4.55))
    add_edge(slide, use_c, ai_c, hw["use_c"], hw["ai_c"], "hasActiveIngredient", (7.35, 3.05))

    # ---------- boxes on top ----------
    add_node_box(slide, *apvma, 2.15, 0.78, "APVMA / PubCRIS", ["official source system"], C_SOURCE)
    add_node_box(slide, *aus, 1.8, 0.78, country["label_en"], ["CountryJurisdiction"], C_COUNTRY)
    add_node_box(slide, *regc, 2.35, 1.0, "Registration 84962",
                 [f"Registered {props.get('registration_date','')} · expires {props.get('expiry_date','')}",
                  "Registered – Current"], C_REG)
    add_node_box(slide, *reg56c, 2.0, 0.95, "Registration 56708",
                 [f"Registered {props56.get('registration_date','')} · expires {props56.get('expiry_date','')}",
                  "sibling registration"], C_REG)
    add_node_box(slide, *fortin, 2.0, 0.78, product["label_en"], ["PesticideProduct"], C_PRODUCT)
    add_node_box(slide, *use_c, 2.0, 0.95, "Approved use",
                 ["official pair asserted", "RegistrationUse"], C_USE)
    add_node_box(slide, *ai_c, 2.0, 0.78, ai["label_en"], ["ActiveIngredientLocal"], C_AI)

    # ---------- right panel ----------
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   ix(8.25), iy(6.7), Inches(4.75 / X0 * SLIDE_W), Inches(5.75 / Y0 * SLIDE_H))
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.fill.background()
    panel.shadow.inherit = False

    add_text(slide, 8.5, 6.27, 4.3, 0.32,
             [[("48 registered crop contexts", 13.5, True, INK)]], align=PP_ALIGN.LEFT)
    add_text(slide, 8.5, 5.99, 4.3, 0.26,
             [[("per the approved use of registration 84962", 9.5, False, MUTED)]], align=PP_ALIGN.LEFT)

    # stacked bar
    bx0, bx1, by, bh = 8.5, 12.75, 5.55, 0.3
    order = sorted(range(len(crops)), key=lambda k: classify(crops[k]))
    counts = [0, 0, 0, 0]
    for k in order:
        counts[classify(crops[k])] += 1
    cx = bx0
    for i, n in enumerate(counts):
        w = (bx1 - bx0) * n / sum(counts)
        seg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ix(cx), iy(by + bh),
                                     Inches(w / X0 * SLIDE_W), Inches(bh / Y0 * SLIDE_H))
        seg.fill.solid()
        seg.fill.fore_color.rgb = CAT_COLORS[i]
        seg.line.color.rgb = WHITE
        seg.line.width = Pt(1.0)
        seg.shadow.inherit = False
        add_text(slide, cx, by, w, bh, [[(str(n), 12, True, WHITE)]])
        cx += w

    # 6 x 8 dot grid
    cols, rows = 8, 6
    gx0, gy0, gdx, gdy = 8.62, 4.85, 0.53, 0.5
    dot_d = 0.26
    pos_of = {}
    for pos, k in enumerate(order):
        pos_of[k] = pos
        r, c = pos // cols, pos % cols
        x, y = gx0 + c * gdx, gy0 - r * gdy
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, ix(x - dot_d / 2), iy(y + dot_d / 2),
                                   Inches(dot_d / X0 * SLIDE_W), Inches(dot_d / Y0 * SLIDE_H))
        d.fill.solid()
        d.fill.fore_color.rgb = CAT_COLORS[classify(crops[k])]
        d.line.color.rgb = WHITE
        d.line.width = Pt(1.0)
        d.shadow.inherit = False

    # representative crop labels (real names on their dots)
    named_crops = {
        "BANANA, OVER 3 YEARS OLD": "Banana",
        "CITRUS OVER 3 YEARS OLD": "Citrus",
        "MANGO, OVER 3 YEARS OLD": "Mango",
        "COTTON": "Cotton",
        "SOYBEAN": "Soybean",
        "SUGAR CANE": "Sugar cane",
        "NAVY BEAN": "Navy bean",
        "PASTURE": "Pasture",
    }
    for full, short in named_crops.items():
        if full not in crops:
            continue
        k = crops.index(full)
        pos = pos_of[k]
        c = pos % cols
        x, y = gx0 + c * gdx, gy0 - (pos // cols) * gdy
        if c <= 4:
            add_text(slide, x + 0.12, y - 0.13, 0.95, 0.26,
                     [[(short, 8.5, False, INK)]], align=PP_ALIGN.LEFT, fill=WHITE)
        else:
            add_text(slide, x - 1.07, y - 0.13, 0.95, 0.26,
                     [[(short, 8.5, False, INK)]], align=PP_ALIGN.RIGHT, fill=WHITE)

    # legend
    lx, ly = 8.55, 1.55
    for i in range(4):
        row, col = i // 2, i % 2
        xx, yy = lx + col * 2.12, ly - row * 0.42
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, ix(xx - 0.09), iy(yy + 0.09),
                                   Inches(0.18 / X0 * SLIDE_W), Inches(0.18 / Y0 * SLIDE_H))
        d.fill.solid()
        d.fill.fore_color.rgb = CAT_COLORS[i]
        d.line.color.rgb = WHITE
        d.line.width = Pt(0.75)
        d.shadow.inherit = False
        add_text(slide, xx + 0.18, yy - 0.14, 1.92, 0.28,
                 [[(f"{CAT_NAMES[i]} ({counts[i]})", 9.5, False, INK)]], align=PP_ALIGN.LEFT)

    # bottom provenance
    add_text(slide, 0.35, 0.08, 11.5, 0.26,
             [[("APVMA PubCRIS  →  PestKG release 2026.08.3_federated    source_record_id retained    official pair asserted",
                9.5, False, MUTED)]], align=PP_ALIGN.LEFT)

    out = OUTPUT / "figure_02_single_site_apvma_editable.pptx"
    prs.save(out)
    print("saved", out)
    print("crops by category:", dict(zip(CAT_NAMES, counts)))


if __name__ == "__main__":
    main()
